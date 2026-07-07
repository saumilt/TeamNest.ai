"""PDF & DOCX export helpers for research, transcripts, summaries, and tasks."""
import io
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

BRAND = colors.HexColor("#facc15")
DARK = colors.HexColor("#0a0a0a")
MUTED = colors.HexColor("#71717a")


def _styles():
    s = getSampleStyleSheet()
    s.add(ParagraphStyle("BrandH1", parent=s["Heading1"], fontSize=22, textColor=DARK, spaceAfter=10))
    s.add(ParagraphStyle("BrandH2", parent=s["Heading2"], fontSize=14, textColor=DARK, spaceAfter=6))
    s.add(ParagraphStyle("Label", parent=s["Normal"], fontSize=8, textColor=MUTED, spaceAfter=3, fontName="Helvetica-Bold"))
    s.add(ParagraphStyle("BodyDark", parent=s["BodyText"], fontSize=10, leading=14, textColor=DARK, spaceAfter=8))
    s.add(ParagraphStyle("Mono", parent=s["BodyText"], fontSize=9, fontName="Courier", textColor=DARK))
    return s


def render_pdf(title: str, sections: list[dict]) -> bytes:
    """Render a structured doc as PDF.
    sections: [{ "label": "...", "body": "..." } | { "label":..., "rows": [["k","v"], ...] } | { "raw": "..." }]
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, topMargin=0.7 * inch, bottomMargin=0.7 * inch,
                            leftMargin=0.8 * inch, rightMargin=0.8 * inch, title=title)
    s = _styles()
    story = []

    # Cover header
    story.append(Paragraph('<font color="#facc15">●</font>  TEAMNEST.AI', s["Label"]))
    story.append(Spacer(1, 4))
    story.append(Paragraph(title, s["BrandH1"]))
    story.append(Paragraph(f'Generated {datetime.utcnow().strftime("%b %d, %Y %H:%M UTC")}', s["Label"]))
    story.append(Spacer(1, 14))

    for sec in sections:
        if "label" in sec:
            story.append(Paragraph(sec["label"].upper(), s["Label"]))
        if "rows" in sec:
            t = Table(sec["rows"], colWidths=[1.6 * inch, 5.0 * inch])
            t.setStyle(TableStyle([
                ("FONT", (0, 0), (-1, -1), "Helvetica", 9),
                ("TEXTCOLOR", (0, 0), (0, -1), MUTED),
                ("TEXTCOLOR", (1, 0), (1, -1), DARK),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e4e4e7")),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))
        if "body" in sec:
            # split paragraphs
            for para in (sec["body"] or "").split("\n\n"):
                para = para.replace("\n", "<br/>")
                story.append(Paragraph(para, s["BodyDark"]))
        if "raw" in sec:
            story.append(Paragraph(sec["raw"], s["Mono"]))
        story.append(Spacer(1, 10))

    # Footer brand bar
    def _on_page(canvas, _doc):
        canvas.saveState()
        canvas.setFillColor(BRAND)
        canvas.rect(0, 0, LETTER[0], 0.18 * inch, fill=1, stroke=0)
        canvas.setFillColor(DARK)
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawString(0.4 * inch, 0.06 * inch, "TEAMNEST.AI")
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(LETTER[0] - 0.4 * inch, 0.06 * inch, f"Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_on_page, onLaterPages=_on_page)
    return buf.getvalue()


def render_docx(title: str, sections: list[dict]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    p = doc.add_paragraph()
    run = p.add_run("● TEAMNEST.AI")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x71, 0x71, 0x7a)

    h = doc.add_heading(title, 0)
    h.runs[0].font.color.rgb = RGBColor(0x0a, 0x0a, 0x0a)

    sub = doc.add_paragraph()
    sub_run = sub.add_run(f'Generated {datetime.utcnow().strftime("%b %d, %Y %H:%M UTC")}')
    sub_run.font.size = Pt(8)
    sub_run.font.color.rgb = RGBColor(0x71, 0x71, 0x7a)

    for sec in sections:
        if "label" in sec:
            doc.add_paragraph().add_run(sec["label"].upper()).bold = True
        if "rows" in sec:
            t = doc.add_table(rows=len(sec["rows"]), cols=2)
            t.autofit = True
            for i, row in enumerate(sec["rows"]):
                t.cell(i, 0).text = str(row[0])
                t.cell(i, 1).text = str(row[1])
                t.cell(i, 0).width = Inches(1.7)
                t.cell(i, 1).width = Inches(4.8)
        if "body" in sec:
            for para in (sec["body"] or "").split("\n\n"):
                doc.add_paragraph(para.replace("\n", " "))
        if "raw" in sec:
            doc.add_paragraph(sec["raw"]).style = doc.styles["Normal"]

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_research(thread: dict, responses: list[dict], saved_title: Optional[str] = None) -> dict:
    """Build sections for an AI research thread export."""
    title = saved_title or (thread.get("question", "AI Research")[:80])
    sections = [
        {"label": "Original Question", "body": thread.get("question") or ""},
        {"label": "Summary",
         "rows": [
             ["Status", thread.get("status", "—")],
             ["Models", ", ".join(thread.get("selected_models") or [])],
             ["Synthesized", "yes" if thread.get("auto_synthesized") else "no"],
             ["Created", thread.get("created_at", "")[:19].replace("T", " ")],
         ]},
        {"label": "Final Answer", "body": thread.get("final_answer") or "—"},
    ]
    for r in responses:
        sections.append({
            "label": f'{r.get("model_name", "Model")} · confidence {r.get("confidence_score", "?")}{" · best" if r.get("selected_as_best") else ""}',
            "body": r.get("answer") or "",
        })
    return {"title": title, "sections": sections}
