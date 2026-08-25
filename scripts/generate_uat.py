"""Generate the TeamNest.ai UAT / Test Case workbook (.xlsx)."""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "/app/frontend/public/TeamNest_UAT_Test_Cases.xlsx"

# (module, feature, title, precondition, steps, test_data, expected, priority, ttype, platform)
CASES = [
    # ---------------- Authentication & Access ----------------
    ("Authentication", "Login", "Valid login", "Registered active user", "1) Go to /login\n2) Enter valid email + password\n3) Click Sign in", "Valid account", "User is signed in and lands on Home (/dashboard).", "P0", "Functional", "Both"),
    ("Authentication", "Login", "Wrong password", "Registered user", "1) Enter correct email, wrong password\n2) Sign in", "wrong password", "Stays on login; toast shows 'Invalid credentials'.", "P0", "Negative", "Both"),
    ("Authentication", "Login", "Unknown email", "-", "1) Enter an email with no account\n2) Sign in", "unknown@x.com", "Toast shows 'Invalid credentials' (no account enumeration).", "P1", "Negative", "Both"),
    ("Authentication", "Login", "Empty fields validation", "-", "1) Leave email/password blank\n2) Sign in", "-", "Form prevents submit / shows required-field validation.", "P2", "Negative", "Both"),
    ("Authentication", "Login throttle", "Lockout after repeated failures", "-", "1) Enter wrong password 5+ times quickly", "wrong x5", "After threshold, HTTP 429 with 'Too many sign-in attempts. Try again in ~N minutes'.", "P1", "Security", "Web"),
    ("Authentication", "Forgot password", "Request reset link", "Registered user", "1) Login -> Forgot password\n2) Enter email\n3) Submit", "valid email", "Generic success message; reset email arrives with a link on the correct domain (teamnest.ai in prod).", "P0", "Functional", "Web"),
    ("Authentication", "Reset password", "Complete reset via link", "Valid reset link (<1h)", "1) Open reset link\n2) Enter a policy-compliant new password\n3) Submit\n4) Login", "New$Pass123", "Password updated; must_change cleared; login with new password succeeds.", "P0", "Functional", "Both"),
    ("Authentication", "Reset password", "Expired link rejected", "Reset link older than 1 hour", "1) Open expired reset link\n2) Try to set password", "-", "Shows 'Invalid or expired reset link'.", "P1", "Negative", "Web"),
    ("Authentication", "Reset password", "Single-use link", "Reset link already used once", "1) Reuse the same link", "-", "Second use is rejected as invalid/expired.", "P1", "Negative", "Web"),
    ("Authentication", "Password policy", "Complexity enforced", "-", "1) On signup/reset/change enter weak password", "abc", "Rejected: needs 8+ chars incl. upper, lower, number & special char.", "P1", "Negative", "Both"),
    ("Authentication", "Temp password", "Forced change on first login", "Provisioned user, must_change_password=true", "1) Login with temp password", "temp pw", "User is forced to set their own password before proceeding.", "P0", "Functional", "Both"),
    ("Authentication", "Temp password", "Expiry", "Temp password older than 7 days", "1) Login with an expired temp password", "old temp pw", "Shows 'Your temporary password has expired. Ask your admin to re-invite you.'", "P1", "Negative", "Web"),
    ("Authentication", "Session", "Logout clears session", "Logged in", "1) Logout\n2) Try to open a protected page", "-", "Session cookie cleared; protected routes redirect to /login.", "P1", "Functional", "Both"),
    ("Authentication", "Session", "Persistence on refresh", "Logged in", "1) Refresh the browser", "-", "User remains logged in (HttpOnly cookie).", "P1", "Functional", "Web"),
    ("Authentication", "MFA", "MFA challenge", "User with MFA enabled", "1) Login with valid creds", "MFA user", "Session is gated behind a passkey/recovery challenge before issuing.", "P1", "Functional", "Web"),

    # ---------------- Signup / Onboarding ----------------
    ("Onboarding", "Signup", "Self-serve workspace creation", "Launch mode allows signup", "1) Go to /signup\n2) Enter name/email/password\n3) Submit", "new user", "Account + new workspace created; user lands on Home.", "P0", "Functional", "Both"),
    ("Onboarding", "Invite code", "Redeem invite code", "Valid active code", "1) Signup with an invite code", "DEVOS100", "User joins the correct program/workspace with expected access.", "P1", "Functional", "Both"),
    ("Onboarding", "Persona pick", "Persona tailors Home", "New user on Home", "1) In 'Make TeamNest yours' pick Just me / Student / A team / Business / Enterprise", "-", "Home quick actions/widgets adjust to the chosen persona.", "P1", "UAT", "Web"),
    ("Onboarding", "Welcome layout picker", "Choose Home layout in welcome flow", "First-time user (or replay welcome)", "1) In welcome card go to 'Choose your Home layout'\n2) Pick ChatGPT/Claude/Chat View", "-", "Selection applies; Home updates live to the chosen layout.", "P1", "UAT", "Web"),
    ("Onboarding", "Welcome tour", "Shows once + skip + replay", "New user", "1) Confirm tour shows once\n2) Skip\n3) Profile -> Show welcome again", "-", "Tour appears once per user; skip dismisses; replay re-opens it.", "P2", "Functional", "Web"),

    # ---------------- Home Layouts ----------------
    ("Home Layouts", "Land on choice", "Login lands on chosen layout", "User with a saved layout", "1) Login", "-", "User lands on /dashboard rendering their chosen layout (not forced to /chats).", "P1", "Functional", "Web"),
    ("Home Layouts", "Change layout", "Switcher lists all looks", "On Home", "1) Click 'Change layout'", "-", "Menu shows Chat View, Start Center, ChatGPT Layout, Claude Layout.", "P1", "UAT", "Web"),
    ("Home Layouts", "Change layout", "Switch to ChatGPT Layout", "On Home", "1) Change layout -> ChatGPT Layout", "-", "Home shows the prompt-first composer ('What can I help with?'); persists after reload.", "P1", "Functional", "Web"),
    ("Home Layouts", "Change layout", "Switch to Claude Layout", "On Home", "1) Change layout -> Claude Layout", "-", "Home shows the calm composer + jump-back-in view; persists.", "P1", "Functional", "Web"),
    ("Home Layouts", "Persistence", "Choice remembered", "-", "1) Pick a layout\n2) Reload / re-login (same browser)", "-", "The chosen layout is retained.", "P2", "Functional", "Web"),
    ("Home Layouts", "Intelligence Banner", "Remembers stats", "Workspace with data", "1) View Home banner", "-", "Shows counts for saved facts, decisions, research threads, documents.", "P2", "UAT", "Web"),
    ("Home Layouts", "Top Action Bar", "Personalized ordering + dismiss", "On main pages", "1) Use several actions\n2) Observe ordering\n3) Dismiss for good", "-", "Actions reorder by usage; 'Dismiss for good' hides the bar persistently.", "P2", "UAT", "Web"),
    ("Home Layouts", "Setup Checklist", "Progress + confetti", "New workspace", "1) Complete checklist steps", "-", "Progress updates; confetti celebration on completion.", "P2", "UAT", "Web"),
    ("Home Layouts", "Show Me How", "Guided walkthrough", "On Home", "1) Trigger 'Show me how'", "-", "Dimmed spotlight walkthrough highlights each step.", "P2", "UAT", "Web"),

    # ---------------- Command Palette ----------------
    ("Command Palette", "Open", "Cmd/Ctrl+K opens palette", "Logged in", "1) Press Cmd/Ctrl+K on any page", "-", "Command palette overlay opens with search input focused.", "P1", "Functional", "Web"),
    ("Command Palette", "Open", "Sidebar Search opens palette", "Logged in", "1) Click 'Search' in sidebar", "-", "Palette opens.", "P2", "Functional", "Web"),
    ("Command Palette", "Search", "Fuzzy filter", "Palette open", "1) Type 'invite'", "invite", "Results narrow to matching action(s), e.g. Invite teammates.", "P1", "Functional", "Web"),
    ("Command Palette", "Quick actions", "Actions navigate", "Palette open", "1) Run New chat / New group / New project / Host meeting / Ask my AI / Upload / Daily standup", "-", "Each routes to the correct page/dialog.", "P1", "Functional", "Web"),
    ("Command Palette", "Navigation", "Go-to items route", "Palette open", "1) Select nav items (Chats, Research, Tasks, etc.)", "-", "Each opens the correct page; role-gated items only for admins/super.", "P1", "Functional", "Web"),
    ("Command Palette", "Data search", "Open research thread", "Has research threads", "1) Type a thread name\n2) Select it", "-", "Navigates to that research thread.", "P2", "Functional", "Web"),
    ("Command Palette", "Data search", "Open project", "Has projects", "1) Type a project name\n2) Select it", "-", "Navigates to /projects/{id}.", "P2", "Functional", "Web"),
    ("Command Palette", "Message person", "Direct chat from palette", "Has teammates", "1) Select a person under 'Message a person'", "-", "Opens an existing direct chat or creates a new one and navigates to it.", "P1", "Functional", "Web"),
    ("Command Palette", "Close", "Esc/backdrop closes + resets", "Palette open with query", "1) Press Esc / click backdrop", "-", "Palette closes; query resets on reopen.", "P2", "Functional", "Web"),

    # ---------------- Chats & Messaging ----------------
    ("Chats", "Direct chat", "Create direct chat", "2 users in workspace", "1) New chat -> pick a person", "-", "Direct chat created and opened.", "P0", "Functional", "Both"),
    ("Chats", "Group chat", "Create group chat", "Logged in", "1) New group -> add members -> create", "-", "Group created with members.", "P0", "Functional", "Both"),
    ("Chats", "Realtime", "Send/receive live", "Two logged-in users", "1) User A sends a message\n2) Observe User B", "-", "Message appears for User B in real time (WebSocket).", "P0", "Functional", "Both"),
    ("Chats", "Invite", "Invite someone new via email", "In a chat/group", "1) Invite someone new -> enter email(s)", "test@x.com", "Invite email sent (correct domain); invitee auto-added when they join.", "P0", "Functional", "Web"),
    ("Chats", "Folders", "Organize chats in folders", "Has chats", "1) Create folder\n2) Move chats in", "-", "Chats grouped under folders.", "P2", "Functional", "Web"),
    ("Chats", "@ai mention", "Model picker", "In a chat", "1) Type @ai", "-", "Picker shows ChatGPT, Claude, Gemini (and others).", "P0", "Functional", "Both"),
    ("Chats", "AI compare", "3 models side-by-side", "In a chat", "1) @ai -> select 3 models -> ask", "-", "All selected models answer in the same thread for comparison.", "P0", "Functional", "Both"),
    ("Chats", "@devmanager", "Summon dev manager", "In a chat", "1) Type @devmanager and a request", "-", "Dev manager AI replies in-thread; can coordinate dev tasks.", "P1", "Functional", "Web"),
    ("Chats", "Attachments", "Attach a file", "In a chat", "1) Attach an image/doc", "-", "File uploads and renders; URL stored (no blobs in DB).", "P1", "Functional", "Both"),
    ("Chats", "Group avatars", "Group avatar displays", "Group chat", "1) Open a group", "-", "Group avatar renders correctly.", "P2", "UI", "Both"),

    # ---------------- AI Research ----------------
    ("AI Research", "Empty state", "Example prompts", "No research yet", "1) Open Research (empty)\n2) Click an example prompt", "-", "Routes to personal AI with the prompt prefilled.", "P1", "UAT", "Web"),
    ("AI Research", "Run research", "Create research thread", "Logged in", "1) Ask AI a research question", "-", "Answer generated; a saved thread appears in Research.", "P0", "Functional", "Both"),
    ("AI Research", "Open thread", "Reopen saved thread", "Has threads", "1) Click a saved thread", "-", "Thread opens with prior Q&A.", "P1", "Functional", "Web"),
    ("AI Research", "Compare", "Multi-model research", "Logged in", "1) Compare models on a research question", "-", "Multiple model answers shown for comparison.", "P1", "Functional", "Web"),

    # ---------------- AI Employees / Marketplace ----------------
    ("AI Employees", "Marketplace", "Browse listings", "Logged in", "1) Open AI Employees / marketplace", "-", "Available AI employees listed.", "P1", "Functional", "Both"),
    ("AI Employees", "Hire", "Install/hire an employee", "Logged in with credits/plan", "1) Hire an AI employee", "-", "License created; employee available in workspace.", "P0", "Functional", "Both"),
    ("AI Employees", "Builder gate", "Non-approved builder blocked", "Non-builder free user", "1) Try to create an AI employee", "buildertest@example.com", "Blocked with 403; prompted to apply to Builder Program.", "P1", "Negative", "Web"),
    ("AI Employees", "Savings", "Savings dashboard", "Has AI usage", "1) Open savings dashboard", "-", "Displays savings metrics without errors.", "P2", "Functional", "Web"),

    # ---------------- Tasks ----------------
    ("Tasks", "Create", "Create + assign task", "Logged in", "1) Create task\n2) Assign + due date", "-", "Task created with assignee and due date.", "P1", "Functional", "Both"),
    ("Tasks", "Complete", "Complete task", "Has a task", "1) Mark task complete", "-", "Task marked done; dashboard counts update.", "P1", "Functional", "Both"),

    # ---------------- Meetings / Calls ----------------
    ("Calls", "Audio", "Start audio call", "LiveKit configured", "1) Start an audio call", "-", "Call connects; audio works.", "P0", "Functional", "Both"),
    ("Calls", "Video", "Start video call", "LiveKit configured", "1) Start a video call", "-", "Call connects; video works.", "P0", "Functional", "Both"),
    ("Calls", "Screen share", "Share screen", "In a call", "1) Start screen share", "-", "Screen is shared to participants.", "P1", "Functional", "Web"),
    ("Calls", "Join", "Second participant joins", "Active call", "1) User B joins the call", "-", "Both participants see/hear each other.", "P0", "Functional", "Both"),
    ("Calls", "Mobile", "Join call on mobile", "Native/dev build", "1) Join a call from the mobile app", "-", "Works on a real device build (NOT Expo Go/web preview).", "P1", "Functional", "Mobile"),

    # ---------------- Knowledge Base / Documents ----------------
    ("Knowledge", "Upload", "Upload a document", "Logged in", "1) Upload pdf/csv/txt", "sample.pdf", "Document uploaded and listed.", "P0", "Functional", "Both"),
    ("Knowledge", "ZIP upload", "Chunked/resumable ZIP", "Logged in", "1) Upload a large ZIP\n2) Interrupt + resume", "big.zip", "Upload chunks; resumes without restarting from scratch.", "P1", "Functional", "Web"),
    ("Knowledge", "Semantic search", "RAG answer with citations", "Docs uploaded", "1) Ask a question about the docs", "-", "Answer synthesized strictly from retrieved excerpts with source citations.", "P0", "Functional", "Both"),
    ("Knowledge", "Ask across docs", "Multi-doc Q&A", "Multiple docs", "1) Ask a question spanning docs", "-", "Relevant sources retrieved and cited.", "P1", "Functional", "Web"),

    # ---------------- AI Memory ----------------
    ("AI Memory", "Save", "Save a fact/decision", "Logged in", "1) Save a fact/decision to memory", "-", "Item stored and visible in AI Memory.", "P1", "Functional", "Both"),
    ("AI Memory", "Retrieve", "Ask my memory", "Has memory items", "1) Ask a question the memory can answer", "-", "Relevant saved item is retrieved in the answer.", "P1", "Functional", "Both"),

    # ---------------- Projects / Folders ----------------
    ("Projects", "Empty state", "Example project chips", "No projects yet", "1) Open Projects (empty)\n2) Click an example chip", "-", "New folder dialog opens PREFILLED with name + description.", "P1", "UAT", "Web"),
    ("Projects", "Create", "Create project folder", "Logged in", "1) New folder -> name + description -> Create", "-", "Folder created and listed.", "P1", "Functional", "Web"),
    ("Projects", "Open", "Open project", "Has projects", "1) Open a project", "-", "Project view shows associated research/tasks/chats.", "P2", "Functional", "Web"),

    # ---------------- Billing / Plans / Credits ----------------
    ("Billing", "Plan view", "See plan + credits", "Logged in", "1) Open Billing", "-", "Current plan and credit balance display.", "P1", "Functional", "Both"),
    ("Billing", "Credit usage", "Credits consumed on AI use", "Has credits", "1) Run an AI action\n2) Check balance", "-", "Balance decreases by expected amount.", "P1", "Functional", "Both"),
    ("Billing", "Governance", "4-tier credit limits", "Governance configured", "1) Exceed a tier limit", "-", "Limits enforced per governance tier.", "P1", "Functional", "Web"),
    ("Billing", "Student plan", "Student plan available", "-", "1) View plans", "-", "Student plan is listed and selectable.", "P2", "Functional", "Both"),
    ("Billing", "Checkout", "Upgrade via Stripe (test)", "Stripe test mode", "1) Upgrade plan\n2) Complete test checkout", "Stripe test card", "Checkout completes; plan upgraded.", "P0", "Functional", "Both"),
    ("Billing", "Top-up", "Buy credits", "Logged in", "1) Buy credits / top-up", "Stripe test card", "Credits added to balance.", "P1", "Functional", "Both"),

    # ---------------- Super Admin ----------------
    ("Super Admin", "Access", "Super-admin only", "Super admin vs normal user", "1) Open /superadmin as each", "-", "Super admin allowed; normal user denied.", "P0", "Security", "Web"),
    ("Super Admin", "Users", "List/search users", "Super admin", "1) Open Users\n2) Search an email", "-", "Users across workspaces listed and searchable.", "P1", "Functional", "Web"),
    ("Super Admin", "Create user", "Provision new user", "Super admin", "1) Add user\n2) Send credentials", "new user", "User created; welcome email link uses the current/custom domain.", "P0", "Functional", "Web"),
    ("Super Admin", "Reset password", "Set/auto-gen password", "Super admin", "1) Reset password (specific + auto)\n2) Use copy button", "-", "Password set; copy button copies exact value; user forced to change on next login.", "P0", "Functional", "Web"),
    ("Super Admin", "Reset link", "Generate reset link", "Super admin", "1) Generate a reset link", "-", "Link generated; domain matches the current/custom domain (teamnest.ai).", "P0", "Functional", "Web"),
    ("Super Admin", "Link Domain Guard", "Warn on wrong-domain link", "Link domain != current domain", "1) Generate a link whose domain differs from current", "-", "Amber warning shows expected vs actual domain + a one-click corrected link.", "P1", "Functional", "Web"),
    ("Super Admin", "Edit user", "Change role/status", "Super admin", "1) Edit a user's role/status\n2) Save", "-", "Changes persist; suspended users cannot log in.", "P1", "Functional", "Web"),
    ("Super Admin", "Delete user", "Delete a user", "Super admin", "1) Delete a user", "-", "User removed; cannot log in afterward.", "P1", "Functional", "Web"),
    ("Super Admin", "Super toggle", "Grant/revoke super admin", "Super admin", "1) Toggle super-admin on another user", "-", "Access granted/revoked accordingly.", "P1", "Security", "Web"),
    ("Super Admin", "Workspaces", "Manage workspaces", "Super admin", "1) Open Workspaces tab", "-", "Workspaces listed and manageable.", "P2", "Functional", "Web"),
    ("Super Admin", "App settings", "Plan credit config", "Super admin", "1) Set free/pro/team monthly credits", "-", "Settings save and apply.", "P2", "Functional", "Web"),

    # ---------------- Integrations / Connectors ----------------
    ("Integrations", "Connectors", "Connect Gmail (OAuth)", "Logged in", "1) Connectors -> connect Gmail", "Google account", "OAuth completes; connector shows connected.", "P1", "Functional", "Web"),
    ("Integrations", "Slack", "Slack messaging", "Slack token provided", "1) Connect Slack\n2) Send/sync a message", "Slack token", "BLOCKED pending updated Slack token/scopes — verify once provided.", "P2", "Functional", "Web"),
    ("Integrations", "HubSpot", "HubSpot CRM", "Private App Token", "1) Connect HubSpot", "HubSpot token", "MOCKED currently — verify against real API once token provided.", "P2", "Functional", "Web"),
    ("Integrations", "Email", "Transactional emails deliver", "Mailgun configured", "1) Trigger reset/invite email", "-", "Email delivered with correct links/domain.", "P1", "Functional", "Web"),

    # ---------------- Marketing Site (public) ----------------
    ("Marketing", "Hero", "Two-line accent headline", "-", "1) Open https://teamnest.ai/", "-", "Hero shows headline with the accent (yellow) second line.", "P2", "UI", "Web"),
    ("Marketing", "Hero proof", "Proof strip stats", "-", "1) View below hero", "-", "Proof strip shows 5+ models / 100% retained / Zero lost / SSO+Audit.", "P2", "UI", "Web"),
    ("Marketing", "Enterprise band", "Request a Demo CTA", "-", "1) Scroll to enterprise band\n2) Click Request a Demo", "-", "Prominent band; CTA routes to demo/support.", "P2", "Functional", "Web"),
    ("Marketing", "Hero A/B", "Alt variant", "-", "1) Open /?hero=alt", "-", "Alt enterprise-focused hero copy renders.", "P3", "Functional", "Web"),
    ("Marketing", "Nav/Pricing", "Marketing nav + pricing", "-", "1) Click Product/Use Cases/Pricing/Business", "-", "All links resolve; pricing renders.", "P2", "Functional", "Web"),
    ("Marketing", "Signup CTA", "Start Free", "-", "1) Click Start Free", "-", "Routes to signup.", "P1", "Functional", "Web"),

    # ---------------- Mobile App (Expo) ----------------
    ("Mobile", "Login", "Mobile login", "Mobile app running", "1) Login on mobile", "valid account", "Signs in and lands on home.", "P0", "Functional", "Mobile"),
    ("Mobile", "Chat", "Mobile chat send/receive", "Logged in on mobile", "1) Send/receive a message", "-", "Real-time messaging works on mobile.", "P0", "Functional", "Mobile"),
    ("Mobile", "Navigation", "Tab navigation", "Logged in on mobile", "1) Switch between tabs", "-", "Tabs navigate; selected tab visually distinct.", "P1", "UI", "Mobile"),
    ("Mobile", "IAP", "RevenueCat purchase", "Native/dev build", "1) Start an in-app purchase", "sandbox account", "Purchase flow works on a real build (NOT Expo Go/web preview).", "P1", "Functional", "Mobile"),
    ("Mobile", "Permissions", "Contextual permission prompts", "Feature needing camera/mic/etc.", "1) Trigger a sensor feature", "-", "Permission requested contextually; denial handled with Open Settings.", "P1", "Functional", "Mobile"),

    # ---------------- Cross-cutting / Non-functional ----------------
    ("Non-functional", "Responsive", "Mobile-web responsive", "-", "1) Resize / open on phone width", "-", "Layout adapts without broken/overlapping elements.", "P1", "UI", "Web"),
    ("Non-functional", "States", "Loading/empty/error states", "-", "1) Visit pages with no data / slow network", "-", "Skeletons on load, helpful empty states, friendly errors + retry.", "P1", "UI", "Both"),
    ("Non-functional", "RBAC", "Role-based access", "member vs admin vs owner", "1) Attempt admin/owner-only actions as a member", "raj@demo.team", "Member is blocked (403) on privileged endpoints.", "P0", "Security", "Both"),
    ("Non-functional", "Performance", "Page load time", "-", "1) Load key pages", "-", "Key pages become interactive within ~3s on a normal connection.", "P2", "Performance", "Both"),
    ("Non-functional", "Security", "No secrets/tokens in client", "-", "1) Inspect network/localStorage", "-", "No auth token in localStorage; session is HttpOnly cookie only.", "P1", "Security", "Web"),
]

HEADERS = [
    "Test ID", "Module", "Feature", "Test Case", "Preconditions", "Test Steps",
    "Test Data", "Expected Result", "Priority", "Type", "Platform",
    "Status", "Actual Result", "Tester", "Date Tested", "Notes / Defect ID",
]

# Colors
BRAND = "111827"      # near-black header
ACCENT = "F5B301"     # TeamNest yellow
GREY = "F3F4F6"
WHITE = "FFFFFF"

thin = Side(style="thin", color="D1D5DB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center")

PRIORITY_FILL = {
    "P0": PatternFill("solid", fgColor="FEE2E2"),
    "P1": PatternFill("solid", fgColor="FEF3C7"),
    "P2": PatternFill("solid", fgColor="E0F2FE"),
    "P3": PatternFill("solid", fgColor="F3F4F6"),
}

wb = Workbook()

# ---------- Instructions sheet ----------
ins = wb.active
ins.title = "Instructions"
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 2
ins.column_dimensions["B"].width = 110

def put(row, text, size=11, bold=False, color="111827"):
    c = ins.cell(row=row, column=2, value=text)
    c.font = Font(size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top")

ins.merge_cells("B2:B2")
put(2, "TeamNest.ai — UAT & Test Cases", size=20, bold=True)
put(3, f"Generated {date.today().isoformat()}  •  {len(CASES)} test cases", size=10, color="6B7280")
put(5, "How to use this workbook", size=13, bold=True, color=ACCENT[:6])
put(6, "1) Go to the 'Test Cases' tab. Each row is one test. Filter by the header dropdowns (Module, Priority, Type, Platform, Status).")
put(7, "2) Assign a tester per Module (or split by Platform: Web / Mobile / Both).")
put(8, "3) For each case: follow the Test Steps, compare against the Expected Result, then set Status = Pass / Fail / Blocked / Not Run.")
put(9, "4) On Fail/Blocked, fill 'Actual Result' and log a defect id in 'Notes / Defect ID'.")
put(10, "5) Run P0 cases first (critical paths), then P1, then P2/P3.")
put(12, "Environments", size=13, bold=True, color=ACCENT[:6])
put(13, "• Production: https://teamnest.ai   • Preview (staging): ask the admin for the current preview URL.")
put(14, "• Test all live-user flows on the intended environment. Reset-link/email tests should use the environment whose domain matches (prod = teamnest.ai).")
put(16, "Suggested test accounts (confirm current values with your admin)", size=13, bold=True, color=ACCENT[:6])
put(17, "• Super admin: sam@funasia.net   • Standard user + member role for permission checks.")
put(18, "• Create a fresh signup to exercise onboarding, empty states, and the welcome layout picker.")
put(20, "Legend", size=13, bold=True, color=ACCENT[:6])
put(21, "Priority: P0 = critical (blockers), P1 = high, P2 = medium, P3 = low.")
put(22, "Type: Functional / UAT / Negative / Security / Performance / UI.  Platform: Web / Mobile / Both.")
put(23, "Note: LiveKit calls, push notifications, and mobile IAP require a real device build — they cannot be validated on Expo Go or the web preview.")

# ---------- Test Cases sheet ----------
ws = wb.create_sheet("Test Cases")
ws.sheet_view.showGridLines = False
ws.append(HEADERS)
for col, _ in enumerate(HEADERS, start=1):
    c = ws.cell(row=1, column=col)
    c.font = Font(bold=True, color=WHITE, size=11)
    c.fill = PatternFill("solid", fgColor=BRAND)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

for i, case in enumerate(CASES, start=1):
    tcid = f"TC-{i:03d}"
    module, feature, title, pre, steps, data, expected, pri, ttype, platform = case
    row = [tcid, module, feature, title, pre, steps, data, expected, pri, ttype, platform, "Not Run", "", "", "", ""]
    ws.append(row)
    r = ws.max_row
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=r, column=col)
        cell.border = border
        cell.alignment = wrap_top
        if r % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=WHITE)
        else:
            cell.fill = PatternFill("solid", fgColor=GREY)
    # priority color
    ws.cell(row=r, column=9).fill = PRIORITY_FILL.get(pri, PatternFill("solid", fgColor=WHITE))
    ws.cell(row=r, column=9).alignment = center
    ws.cell(row=r, column=11).alignment = center  # platform
    ws.cell(row=r, column=12).alignment = center  # status

widths = [9, 15, 16, 30, 24, 40, 16, 44, 8, 12, 9, 10, 30, 12, 14, 22]
for idx, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

# Data validation dropdowns
status_dv = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked"', allow_blank=True)
ws.add_data_validation(status_dv)
status_dv.add(f"L2:L{ws.max_row}")
pri_dv = DataValidation(type="list", formula1='"P0,P1,P2,P3"', allow_blank=True)
ws.add_data_validation(pri_dv)
pri_dv.add(f"I2:I{ws.max_row}")

# ---------- Summary sheet ----------
sm = wb.create_sheet("Summary")
sm.sheet_view.showGridLines = False
sm.column_dimensions["A"].width = 22
sm.column_dimensions["B"].width = 12
sm.column_dimensions["D"].width = 14
sm.column_dimensions["E"].width = 12

def hdr(cell, text):
    c = sm[cell]; c.value = text
    c.font = Font(bold=True, color=WHITE); c.fill = PatternFill("solid", fgColor=BRAND)
    c.alignment = center; c.border = border

from collections import Counter
by_module = Counter(c[0] for c in CASES)
by_pri = Counter(c[7] for c in CASES)
by_plat = Counter(c[9] for c in CASES)

sm["A1"] = "TeamNest UAT — Summary"; sm["A1"].font = Font(bold=True, size=16)
sm["A2"] = f"Total test cases: {len(CASES)}"; sm["A2"].font = Font(size=11, color="6B7280")

hdr("A4", "Module"); hdr("B4", "Count")
r = 5
for m, n in sorted(by_module.items()):
    sm.cell(row=r, column=1, value=m).border = border
    sm.cell(row=r, column=2, value=n).border = border
    sm.cell(row=r, column=2).alignment = center
    r += 1

hdr("D4", "Priority"); hdr("E4", "Count")
r = 5
for p in ["P0", "P1", "P2", "P3"]:
    if by_pri.get(p):
        sm.cell(row=r, column=4, value=p).border = border
        sm.cell(row=r, column=5, value=by_pri[p]).border = border
        sm.cell(row=r, column=5).alignment = center
        r += 1
r += 1
hdr(f"D{r}", "Platform"); hdr(f"E{r}", "Count")
r += 1
for pl, n in sorted(by_plat.items()):
    sm.cell(row=r, column=4, value=pl).border = border
    sm.cell(row=r, column=5, value=n).border = border
    sm.cell(row=r, column=5).alignment = center
    r += 1

wb.save(OUT)
print("Wrote", OUT, "with", len(CASES), "cases")
